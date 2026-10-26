from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib
import json
import os
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from progress import emit_progress
from testdata_common import INTERNAL_KEY
from validate_testdata_workbook import load_metadata, validate


DELEGATED_AUTH_BLOCKERS = (
    "CLIENT_ID",
    "CLIENT_SECRET",
    "CLIENT_CERTIFICATE_PATH",
    "CLIENT_CERTIFICATE_PASSWORD",
)


def require_delegated_auth() -> None:
    configured = [name for name in DELEGATED_AUTH_BLOCKERS if os.environ.get(name)]
    if configured:
        raise RuntimeError(
            "This workflow only accepts user-delegated authentication configured by "
            f"the Microsoft dv-connect skill. Remove: {', '.join(configured)}."
        )


def import_get_client():
    candidates = [
        Path(os.environ["DV_TESTDATA_AUTH_DIR"]) if os.environ.get("DV_TESTDATA_AUTH_DIR") else None,
        Path.cwd() / "scripts",
        Path.home() / ".copilot" / "installed-plugins" / "awesome-copilot" / "dataverse" / "scripts",
    ]
    for candidate in candidates:
        if candidate and (candidate / "auth.py").exists():
            sys.path.insert(0, str(candidate))
            return importlib.import_module("auth").get_client
    raise RuntimeError(
        "auth.py was not found. Set DV_TESTDATA_AUTH_DIR to the Dataverse plugin scripts folder."
    )


def serializable(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return value


def chunks(items: list[Any], size: int):
    for start in range(0, len(items), size):
        yield items[start:start + size]


def write_manifest(manifest: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(manifest, handle, indent=2)


def workbook_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def is_multiple_unsupported(error: Exception, operation: str) -> bool:
    message = str(error).casefold()
    return (
        operation.casefold() in message
        and ("does not support" in message or "messageprocessor.empty" in message)
    )


def is_create_multiple_unsupported(error: Exception) -> bool:
    return is_multiple_unsupported(error, "CreateMultiple")


def normalize_created_ids(result: Any) -> list[str]:
    if isinstance(result, str):
        return [result]
    return [str(value) for value in result]


def resumed_record_map(
    table_name: str,
    items: list[dict[str, str]],
    approved_row_keys: set[str],
) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in items:
        if not isinstance(item, dict):
            raise RuntimeError(
                f"{table_name}: resume manifest uses the legacy GUID-only format and "
                "cannot safely resume individual workbook rows."
            )
        row_key = str(item.get("row_key") or "")
        record_id = str(item.get("id") or "")
        if not row_key or not record_id:
            raise RuntimeError(f"{table_name}: resume manifest has an incomplete record entry.")
        if row_key not in approved_row_keys:
            raise RuntimeError(
                f"{table_name}: resume manifest contains unknown row key {row_key!r}."
            )
        if row_key in result:
            raise RuntimeError(
                f"{table_name}: resume manifest contains duplicate row key {row_key!r}."
            )
        result[row_key] = record_id
    return result


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Apply an approved Dataverse test-data Excel workbook."
    )
    parser.add_argument("--workbook", required=True)
    parser.add_argument(
        "--environment-url",
        required=True,
        help="Environment explicitly selected through the Microsoft dv-connect skill.",
    )
    parser.add_argument("--seed-email", required=True)
    parser.add_argument(
        "--approved",
        required=True,
        choices=["YES"],
        help="Explicit approval gate. Must be exactly YES.",
    )
    parser.add_argument("--chunk-size", type=int, default=500)
    parser.add_argument("--manifest-dir", default="testdata-manifests")
    parser.add_argument(
        "--resume-manifest",
        help="Manifest of an interrupted run; tables already in it are reused, not recreated.",
    )
    args = parser.parse_args()

    emit_progress(
        8,
        "running",
        "Revalidating the approved workbook and connection before the first Dataverse write.",
        "Create one retained probe per table, then create the remaining rows in dependency order.",
    )
    os.environ["DATAVERSE_URL"] = args.environment_url.rstrip("/")
    workbook_path = Path(args.workbook)
    rows_by_table, errors = validate(workbook_path)
    if errors:
        emit_progress(
            8,
            "blocked",
            f"The approved workbook has {len(errors)} validation error(s); no records were created.",
            "Correct and revalidate the workbook, then request approval again.",
        )
        for error in errors:
            print(f"ERROR: {error}")
        raise SystemExit("Workbook validation failed; no records were created.")

    workbook = load_workbook(workbook_path, data_only=False)
    run, tables = load_metadata(workbook)
    get_client = import_get_client()
    require_delegated_auth()
    client = get_client("dv-data")
    active_environment = (os.environ.get("DATAVERSE_URL") or "").rstrip("/")
    planned_environment = (run.get("environment") or "").rstrip("/")
    if not active_environment or active_environment.casefold() != planned_environment.casefold():
        emit_progress(
            8,
            "blocked",
            "The active Dataverse environment does not match the approved workbook.",
            "Reconnect to the workbook environment or generate a new preview for the active environment.",
        )
        raise RuntimeError(
            "The active Dataverse environment does not match the workbook. "
            f"Workbook={planned_environment!r}, active={active_environment!r}. No records were created."
        )
    if args.environment_url.rstrip("/").casefold() != planned_environment.casefold():
        emit_progress(
            8,
            "blocked",
            "The explicitly selected environment does not match the approved workbook.",
            "Select the workbook environment or generate and approve a new workbook.",
        )
        raise RuntimeError(
            "The selected environment does not match the workbook environment. "
            "No records were created."
        )
    if args.seed_email.strip().casefold() != str(run.get("seed_email") or "").casefold():
        emit_progress(
            8,
            "blocked",
            "The supplied seed email does not match the approved workbook.",
            "Use the workbook seed email or generate and approve a new preview.",
        )
        raise RuntimeError(
            "The supplied seed email does not match the approved workbook. "
            "No records were created."
        )

    id_by_key: dict[str, dict[str, str]] = {}
    digest = workbook_hash(workbook_path)
    manifest = {
        "environment": run.get("environment"),
        "created_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "workbook": str(workbook_path.resolve()),
        "workbook_sha256": digest,
        "order": run["order"],
        "records": {},
        "actions": {},
        "deferred_updates": {},
    }
    manifest_path = (
        Path(args.resume_manifest)
        if args.resume_manifest
        else Path(args.manifest_dir)
        / f"{run.get('name', 'testdata')}-{run['order'][0]}-{run.get('run_id', 'run')}.json"
    )

    resumed: dict[str, list[dict[str, str]]] = {}
    if args.resume_manifest:
        with Path(args.resume_manifest).open(encoding="utf-8") as handle:
            previous = json.load(handle)
        if str(previous.get("environment") or "").rstrip("/").casefold() != planned_environment.casefold():
            raise RuntimeError(
                "Resume manifest belongs to a different environment. No records were created."
            )
        if str(previous.get("workbook")) != str(workbook_path.resolve()):
            raise RuntimeError("Resume manifest belongs to a different workbook. No records were created.")
        if previous.get("workbook_sha256") and previous["workbook_sha256"] != digest:
            raise RuntimeError(
                "Resume manifest belongs to a different workbook revision. No records were created."
            )
        resumed = previous.get("records") or {}
        manifest = previous
        manifest.setdefault("workbook_sha256", digest)
        manifest.setdefault("order", run["order"])
        manifest.setdefault("actions", {})
        manifest.setdefault("deferred_updates", {})

    for table_name in run["order"]:
        approved_row_keys = {
            str(row[INTERNAL_KEY]) for row in rows_by_table[table_name]
        }
        completed = resumed_record_map(
            table_name,
            resumed.get(table_name, []),
            approved_row_keys,
        )
        manifest["records"].setdefault(table_name, list(resumed.get(table_name, [])))
        if len(completed) == len(rows_by_table[table_name]):
            id_by_key[table_name] = completed
            emit_progress(
                8,
                "running",
                f"{table_name}: reusing {len(completed)} previously created record(s) "
                "from the resume manifest.",
                "Continue with the next dependency-ordered table.",
            )
        table = tables[table_name]
        payloads: list[dict[str, Any]] = []
        keys: list[str] = []
        lookup_columns = {lookup["column"]: lookup for lookup in table.get("lookups", [])}
        allowed_columns = {column["logical_name"] for column in table.get("columns", [])}

        for row in rows_by_table[table_name]:
            row_key = str(row[INTERNAL_KEY])
            if row_key in completed:
                continue
            payload = {
                key: serializable(value)
                for key, value in row.items()
                if key in allowed_columns and value is not None and value != ""
            }
            for column_name, lookup in lookup_columns.items():
                if lookup.get("deferred"):
                    continue
                target_key = row.get(column_name)
                if not target_key:
                    continue
                if str(target_key) in (lookup.get("existing_ids") or []):
                    target_guid = str(target_key)
                else:
                    target_guid = id_by_key[lookup["target_table"]][str(target_key)]
                if lookup["target_table"] in tables:
                    target_entity_set = tables[lookup["target_table"]]["entity_set_name"]
                else:
                    info = client.tables.get(lookup["target_table"])
                    target_entity_set = info.get("EntitySetName") or info.get("entity_set_name")
                payload[f"{lookup['navigation_property']}@odata.bind"] = (
                    f"/{target_entity_set}({target_guid})"
                )
            payloads.append(payload)
            keys.append(row_key)

        dependencies = [
            lookup["target_table"]
            for lookup in table.get("lookups", [])
            if not lookup.get("existing_ids") and not lookup.get("deferred")
        ]
        emit_progress(
            8,
            "running",
            f"{table_name}: preparing {len(payloads)} remaining record(s). "
            + (
                "Generated parent records are already available from: "
                + ", ".join(sorted(set(dependencies)))
                + "."
                if dependencies
                else "This table has no generated-parent dependency."
            ),
            "Create the first approved row as a retained probe before sending the remaining chunks.",
        )

        def record_created(row_key: str, record_id: str) -> None:
            completed[row_key] = record_id
            manifest["records"][table_name].append(
                {"row_key": row_key, "id": record_id}
            )
            write_manifest(manifest, manifest_path)

        if payloads:
            try:
                probe_ids = normalize_created_ids(
                    client.records.create(table_name, payloads[0])
                )
            except Exception as error:
                emit_progress(
                    8,
                    "blocked",
                    f"{table_name}: Dataverse rejected the current probe or chunk: {error}",
                    "Review the business-rule error; regenerate and reapprove if values must change.",
                )
                raise
            if len(probe_ids) != 1:
                raise RuntimeError(
                    f"{table_name}: Dataverse returned {len(probe_ids)} IDs for the probe row."
                )
            record_created(keys[0], probe_ids[0])
            emit_progress(
                8,
                "running",
                f"{table_name}: created {len(completed)}/{len(rows_by_table[table_name])} record(s).",
                "Continue the current table or advance to the next dependent table.",
            )

        remaining_payloads = payloads[1:]
        remaining_keys = keys[1:]
        offset = 0
        for payload_chunk in chunks(remaining_payloads, args.chunk_size):
            chunk_keys = remaining_keys[offset:offset + len(payload_chunk)]
            try:
                if len(payload_chunk) == 1:
                    result_ids = [
                        str(client.records.create(table_name, payload_chunk[0]))
                    ]
                else:
                    result_ids = normalize_created_ids(
                        client.records.create(table_name, payload_chunk)
                    )
            except Exception as error:
                if not is_create_multiple_unsupported(error):
                    emit_progress(
                        8,
                        "blocked",
                        f"{table_name}: Dataverse rejected the current chunk: {error}",
                        "Review the business-rule error; regenerate and reapprove if values must change.",
                    )
                    raise
                result_ids = []
                for row_key, payload in zip(chunk_keys, payload_chunk):
                    record_id = str(client.records.create(table_name, payload))
                    result_ids.append(record_id)
                    record_created(row_key, record_id)
                    emit_progress(
                        8,
                        "running",
                        f"{table_name}: CreateMultiple is unavailable; created "
                        f"{len(completed)}/{len(rows_by_table[table_name])} record(s) individually.",
                        "Continue the sequential fallback for this table.",
                    )
            else:
                if len(result_ids) != len(chunk_keys):
                    raise RuntimeError(
                        f"{table_name}: Dataverse returned {len(result_ids)} IDs for "
                        f"{len(chunk_keys)} rows in a chunk."
                    )
                for row_key, record_id in zip(chunk_keys, result_ids):
                    record_created(row_key, record_id)
            offset += len(payload_chunk)
            emit_progress(
                8,
                "running",
                f"{table_name}: created {len(completed)}/{len(rows_by_table[table_name])} record(s).",
                "Continue the current table or advance to the next dependent table.",
            )

        if len(completed) != len(rows_by_table[table_name]):
            raise RuntimeError(
                f"{table_name}: manifest contains {len(completed)} IDs for "
                f"{len(rows_by_table[table_name])} rows."
            )
        id_by_key[table_name] = completed

        completed_actions = set(manifest["actions"].get(table_name, []))
        for action_index, action in enumerate(table.get("post_create_actions", [])):
            action_key = f"{action_index}:{action.get('kind')}"
            if action_key in completed_actions:
                continue
            if action.get("kind") != "set_state":
                raise RuntimeError(
                    f"{table_name}: unsupported approved post-create action {action.get('kind')!r}."
                )
            try:
                client.records.update(
                    table_name,
                    list(completed.values()),
                    action["data"],
                )
            except Exception as error:
                if not is_multiple_unsupported(error, "UpdateMultiple"):
                    raise
                for record_id in completed.values():
                    client.records.update(table_name, record_id, action["data"])
            manifest["actions"].setdefault(table_name, []).append(action_key)
            write_manifest(manifest, manifest_path)
            emit_progress(
                8,
                "running",
                f"{table_name}: completed approved post-create action: "
                f"{action.get('description') or action.get('kind')}.",
                "Continue with dependent tables.",
            )

    for table_name in run["order"]:
        table = tables[table_name]
        rows = rows_by_table[table_name]
        for lookup in table.get("lookups", []):
            if not lookup.get("deferred"):
                continue
            update_key = f"{table_name}.{lookup['column']}"
            completed_updates = set(manifest["deferred_updates"].get(update_key, []))
            for row in rows:
                row_key = str(row[INTERNAL_KEY])
                if row_key in completed_updates:
                    continue
                target_key = str(row[lookup["column"]])
                if target_key in (lookup.get("existing_ids") or []):
                    target_guid = target_key
                else:
                    target_guid = id_by_key[lookup["target_table"]][target_key]
                if lookup["target_table"] in tables:
                    target_entity_set = tables[lookup["target_table"]]["entity_set_name"]
                else:
                    info = client.tables.get(lookup["target_table"])
                    target_entity_set = info.get("EntitySetName") or info.get("entity_set_name")
                client.records.update(
                    table_name,
                    id_by_key[table_name][row_key],
                    {
                        f"{lookup['navigation_property']}@odata.bind": (
                            f"/{target_entity_set}({target_guid})"
                        )
                    },
                )
                manifest["deferred_updates"].setdefault(update_key, []).append(row_key)
                write_manifest(manifest, manifest_path)

    emit_progress(
        8,
        "complete",
        "All approved records were created and recorded in the rollback manifest.",
        f"Use {manifest_path.resolve()} if this run must be reviewed or rolled back.",
    )
    print(f"Manifest written: {manifest_path.resolve()}")


if __name__ == "__main__":
    main()
